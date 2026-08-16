"""Failing tests for T-063 (module layer) — EIA-860M generator inventory.

``parse_860m(xlsx_bytes: bytes, sheet: str) -> list[dict]`` (pure) — the
monthly EIA-860M workbook's "Operating" or "Planned" sheet -> normalized
rows {plant, entity, county, state, mw, technology, sector, status,
planned_online} (all str except mw float). The real workbook (verified
against june_generator2026.xlsx, fetched 2026-08-14) carries a title row,
a blank row, the header on ROW 3, data rows, then a trailing blank row
and a "NOTES:" footer row — footer/blank rows (no Plant Name) are
skipped. status is derived from the SHEET ("operating"/"planned"), never
from the sheet's own Status column. planned_online is "" on Operating
rows and "YYYY-MM" on Planned rows only when BOTH Planned Operation
Month and Year are present (blanks are literal " " cells in the real
file). Unknown sheet names, a workbook missing the requested sheet,
bytes that are not an xlsx, or a sheet whose header row cannot be found
are LOUD ValueErrors — a silent [] on a schema change would quietly zero
the dataset.

``latest_links(html: str) -> list[str]`` (pure) — the eia860m index page
HTML -> absolute xlsx URLs matching <month>_generator<year>.xlsx, newest
first by (year, month), deduplicated. The live page also lists dead
FUTURE-month placeholder links (301 -> an HTML landing page), so the
fetcher must verify it received a real xlsx and fall through to the
next-newest link; ordering here is pure and keeps placeholders in place.

Jurisdiction join: the SAME counties join through the SAME table —
eia860m re-exports iso_queues.join_jurisdictions/load_mapping (identity,
not a copy), keyed on f"{county}, {state}" exactly (T-059 honesty pin:
misses returned, never fuzzy-matched or dropped).
"""

import io

import openpyxl
import pytest

try:
    from onrecord.ingest import eia860m
except Exception:  # pragma: no cover - red phase
    eia860m = None

ROW_KEYS = {
    "plant", "entity", "county", "state", "mw", "technology", "sector", "status",
    "planned_online",
}


def _attr(name):
    if eia860m is None or not hasattr(eia860m, name):
        pytest.fail(f"onrecord.ingest.eia860m.{name} does not exist yet (T-063 red)")
    return getattr(eia860m, name)


# Headers and data rows below are copied from the real june_generator2026.xlsx
# (header row 3 verified; Operating header truncated after its real 23rd
# column "Status" — the parser must key columns by NAME, not position).

OPERATING_HEADER = [
    "Entity ID", "Entity Name", "Plant ID", "Plant Name", "Google Map", "Bing Map",
    "Plant State", "County", "Balancing Authority Code", "Sector", "Generator ID",
    "Unit Code", "Nameplate Capacity (MW)", "Net Summer Capacity (MW)",
    "Net Winter Capacity (MW)", "Technology", "Energy Source Code", "Prime Mover Code",
    "Operating Month", "Operating Year", "Planned Retirement Month",
    "Planned Retirement Year", "Status",
]

PLANNED_HEADER = [
    "Entity ID", "Entity Name", "Plant ID", "Plant Name", "Google Map", "Bing Map",
    "Plant State", "County", "Balancing Authority Code", "Sector", "Generator ID",
    "Unit Code", "Nameplate Capacity (MW)", "Net Summer Capacity (MW)",
    "Net Winter Capacity (MW)", "Technology", "Energy Source Code", "Prime Mover Code",
    "Planned Operation Month", "Planned Operation Year", "Status", "Latitude",
    "Longitude",
]

# Real Operating rows: Sand Point appears twice with different Status codes —
# both normalize to status "operating" (sheet-derived, column ignored).
SAND_POINT_SB = [
    63560, "Sand Point Generating, LLC", 1, "Sand Point", "Map", "Map", "AK",
    "Aleutians East", "", "Electric Utility", "1", "", 0.9, 0.4, 0.4,
    "Petroleum Liquids", "DFO", "IC", 12, 2000, " ", " ",
    "(SB) Standby/Backup: available for service but not normally used",
]
SAND_POINT_OP = [
    63560, "Sand Point Generating, LLC", 1, "Sand Point", "Map", "Map", "AK",
    "Aleutians East", "", "Electric Utility", "2", "", 0.9, 0.4, 0.4,
    "Petroleum Liquids", "DFO", "IC", 12, 2000, " ", " ", "(OP) Operating",
]
RAGSDALE_MS = [
    65348, "Ragsdale Solar, LLC", 66240, "Ragsdale Solar LLC", "Map", "Map", "MS",
    "Madison", "MISO", "IPP Non-CHP", "GEN01", "", 100, 100, 100,
    "Solar Photovoltaic", "SUN", "PV", 11, 2024, " ", " ", "(OP) Operating",
]

# Real Planned rows: month/year blanks are literal " " cells; present values
# are ints. No real row mixes one present with one blank, so GAIA_NO_YEAR is
# a shape-faithful variant of the real Gaia row pinning the BOTH-required rule.
GAIA_TX = [
    65815, "Sunraycer Assets I LLC", 66342, "Gaia Hybrid", "Map", "Map", "TX",
    "Navarro", "ERCO", "IPP Non-CHP", "GAIAS", "", 152.7, 152.7, 152.7,
    "Solar Photovoltaic", "SUN", "PV", " ", " ", "(OP) Operating",
    32.160767, -96.19527,
]
STURBRIDGE_MA = [
    65043, "Madison Energy Investments LLC", 63677, "Sturbridge Road Solar", "Map",
    "Map", "MA", "Worcester", "ISNE", "IPP Non-CHP", "STURB", "", 5, 5, 5,
    "Solar Photovoltaic", "SUN", "PV", 6, 2026,
    "(TS) Construction complete, but not yet in commercial operation",
    42.174611, -72.01786,
]
GAIA_NO_YEAR = GAIA_TX[:18] + [6, " "] + GAIA_TX[20:]

NOTES_TEXT = (
    "NOTES:\nCapacity from facilities with a total generator nameplate capacity "
    "less than 1 MW are excluded from this report."
)


def _xlsx(sheet, title, header, rows):
    """Real workbook layout: title, blank, header on row 3, data, then the
    real trailing blank row and NOTES footer row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append([title])
    ws.append(["" for _ in header])
    ws.append(header)
    for row in rows:
        ws.append(row)
    ws.append([""])
    ws.append([NOTES_TEXT])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _operating_bytes():
    return _xlsx(
        "Operating", "Inventory of Operating Generators as of June 2026",
        OPERATING_HEADER, [SAND_POINT_SB, SAND_POINT_OP, RAGSDALE_MS],
    )


def _planned_bytes():
    return _xlsx(
        "Planned", "Inventory of Planned Generators as of June 2026",
        PLANNED_HEADER, [GAIA_TX, STURBRIDGE_MA, GAIA_NO_YEAR],
    )


def test_parse_operating_normalized_rows():
    rows = _attr("parse_860m")(_operating_bytes(), "Operating")
    assert len(rows) == 3  # trailing blank + NOTES footer rows skipped
    assert rows[0] == {
        "plant": "Sand Point", "entity": "Sand Point Generating, LLC",
        "county": "Aleutians East", "state": "AK", "mw": 0.9,
        "technology": "Petroleum Liquids", "sector": "Electric Utility",
        "status": "operating", "planned_online": "",
    }
    # status comes from the sheet, not the Status column
    assert {r["status"] for r in rows} == {"operating"}
    assert (rows[2]["county"], rows[2]["state"], rows[2]["mw"]) == ("Madison", "MS", 100.0)


def test_parse_planned_planned_online():
    rows = _attr("parse_860m")(_planned_bytes(), "Planned")
    assert len(rows) == 3
    assert {r["status"] for r in rows} == {"planned"}
    assert rows[1]["plant"] == "Sturbridge Road Solar"
    assert rows[1]["planned_online"] == "2026-06"  # month 6 zero-padded
    assert rows[0]["planned_online"] == ""  # both cells are blank " " (real Gaia row)
    assert rows[2]["planned_online"] == ""  # month present, year blank: BOTH required
    assert rows[1]["mw"] == 5.0


def test_row_shape_and_types():
    op = _attr("parse_860m")(_operating_bytes(), "Operating")
    pl = _attr("parse_860m")(_planned_bytes(), "Planned")
    assert op and pl
    for row in op + pl:
        assert set(row) == ROW_KEYS
        assert isinstance(row["mw"], float)
        for key in ROW_KEYS - {"mw"}:
            assert isinstance(row[key], str)


def test_unknown_sheet_loud():
    parse = _attr("parse_860m")
    content = _operating_bytes()
    with pytest.raises(ValueError):
        parse(content, "Retired")  # a real sheet, but not one we normalize
    with pytest.raises(ValueError):
        parse(content, "operating")  # exact sheet names only
    with pytest.raises(ValueError):
        parse(_planned_bytes(), "Operating")  # workbook lacks the sheet


def test_not_an_xlsx_loud():
    with pytest.raises(ValueError):
        _attr("parse_860m")(b"not a zip archive", "Operating")


def test_missing_header_row_loud():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Operating"
    ws.append(["Inventory of Operating Generators as of June 2026"])
    ws.append(["nothing", "like", "the", "real", "header"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        _attr("parse_860m")(buf.getvalue(), "Operating")


def test_join_is_reused_from_iso_queues():
    from onrecord.ingest import iso_queues

    assert _attr("join_jurisdictions") is iso_queues.join_jurisdictions
    assert _attr("load_mapping") is iso_queues.load_mapping


def test_join_jurisdictions_on_860m_rows():
    rows = _attr("parse_860m")(_operating_bytes(), "Operating")
    mapping = {"Madison, MS": "Madison County, MS"}
    hits, misses = _attr("join_jurisdictions")(rows, mapping)
    assert [h["jurisdiction"] for h in hits] == ["Madison County, MS"]
    assert hits[0]["plant"] == "Ragsdale Solar LLC"
    assert misses == [rows[0], rows[1]]  # unchanged, in order
    assert "jurisdiction" not in rows[2]  # inputs never mutated


# Real hrefs from https://www.eia.gov/electricity/data/eia860m/ (2026-08-15):
# the current file lives under /xls/, older months under /archive/xls/, and
# future placeholder months (october/december 2026) are listed but dead.
INDEX_HTML = (
    '<a href="/electricity/data/eia860m/xls/december_generator2026.xlsx">re</a>\n'
    '<a href="/electricity/data/eia860m/xls/june_generator2026.xlsx">June</a>\n'
    '<a href="/electricity/data/eia860m/archive/xls/may_generator2026.xlsx">May</a>\n'
    '<a href="/electricity/data/eia860m/archive/xls/december_generator2025.xlsx">D</a>\n'
    '<a href="/electricity/data/eia860m/archive/xls/august_generator2015.xlsx">A</a>\n'
    '<a href="/electricity/data/eia860m/xls/june_generator2026.xlsx">dupe</a>\n'
    '<a href="/electricity/data/860m_faq.pdf">not a generator xlsx</a>\n'
)


def test_latest_links_newest_first():
    links = _attr("latest_links")(INDEX_HTML)
    assert links == [
        "https://www.eia.gov/electricity/data/eia860m/xls/december_generator2026.xlsx",
        "https://www.eia.gov/electricity/data/eia860m/xls/june_generator2026.xlsx",
        "https://www.eia.gov/electricity/data/eia860m/archive/xls/may_generator2026.xlsx",
        "https://www.eia.gov/electricity/data/eia860m/archive/xls/december_generator2025.xlsx",
        "https://www.eia.gov/electricity/data/eia860m/archive/xls/august_generator2015.xlsx",
    ]  # deduped; the dead future-month link stays first — the FETCHER skips non-xlsx


def test_latest_links_no_matches():
    assert _attr("latest_links")("<html><body>no spreadsheets here</body></html>") == []
